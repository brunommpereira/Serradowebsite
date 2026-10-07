"""Consolida a base de dados de atletas do Serrado FC com os resultados do
Troféu Almada em Atletismo num único ficheiro Excel (+ CSV para importar no backend).

Uso:
  python3 consolidate.py --bd BD.xlsx --form Respostas.xlsx [--form Respostas2.xlsx] \
      --results out/resultados.csv --out out/serrado-consolidado.xlsx [--import-dir out/import]

ATENÇÃO: os ficheiros de entrada e de saída contêm dados pessoais — nunca os
coloques no repositório (a pasta out/ está no .gitignore).
"""
import argparse, csv, os, re, sys, unicodedata, warnings, datetime as dt
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')
ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
ap.add_argument('--bd', required=True, help='Excel da base de dados (lista mestra + folhas «SFC | OTC …»)')
ap.add_argument('--form', action='append', default=[], help='Excel de respostas do formulário de inscrição (repetível)')
ap.add_argument('--results', required=True, help='CSV normalizado (saída de normalize.py)')
ap.add_argument('--out', required=True, help='Excel consolidado a gerar')
ap.add_argument('--import-dir', help='Pasta para os CSV de importação no backend (athletes, results)')
ap.add_argument('--regs', default=os.path.join(os.path.dirname(__file__), 'regulamentos.tsv'))
args = ap.parse_args()
BD, RESULTS, OUT = args.bd, args.results, args.out
REF_DATE = dt.date.today()

# ---------------------------------------------------------------- utilidades
def norm(s):
    s = unicodedata.normalize('NFD', str(s or '').upper())
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Z ]+', ' ', s).split()

STOP = {'DE', 'DA', 'DO', 'DAS', 'DOS', 'E'}
def tokens(s):
    return [t for t in norm(s) if t not in STOP and len(t) > 1]

def digits(v):
    if v in (None, ''):
        return ''
    if isinstance(v, float):
        v = int(v)
    return re.sub(r'\D', '', str(v))

def cc_norm(v):
    """N.º de CC/BI base: sem zeros à esquerda nem dígito/letras de controlo."""
    d = digits(v).lstrip('0')
    return d[:8]

def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None

def rows_of(path, sheet=None, idx=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[idx]
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else '' for h in next(it)]
    out = []
    for r in it:
        if not any(c not in (None, '') for c in r):
            continue
        out.append({h: r[i] if i < len(r) else None for i, h in enumerate(head) if h})
    return out

def tok_eq(x, y):
    """Igual, abreviado (PREG/PREGO) ou com uma letra trocada/em falta (BUNO/BRUNO)."""
    if x == y:
        return True
    if min(len(x), len(y)) < 4:
        return False
    if x.startswith(y) or y.startswith(x):
        return True
    if abs(len(x) - len(y)) > 1:
        return False
    # distância de edição <= 1
    i = 0
    while i < min(len(x), len(y)) and x[i] == y[i]:
        i += 1
    return x[i + 1:] == y[i + 1:] or x[i + 1:] == y[i:] or x[i:] == y[i + 1:]

def name_match(a, b):
    """Todos os tokens do nome mais curto existem no mais longo (1.º nome igual)."""
    ta, tb = tokens(a), tokens(b)
    if not ta or not tb or not tok_eq(ta[0], tb[0]):
        return False
    short, long_ = (ta, tb) if len(ta) <= len(tb) else (tb, ta)
    return all(any(tok_eq(t, u) for u in long_) for t in short)

# ------------------------------------------------- 1. base de dados de atletas
athletes = {}          # chave (CC ou nome+data) -> registo
def key_for(cc, name, birth):
    return f'CC{cc}' if cc else f'N{" ".join(tokens(name))}|{birth}'

def upsert(src, rec):
    k = key_for(rec.get('CC'), rec.get('Nome'), rec.get('Data nascimento'))
    a = athletes.setdefault(k, {'Fontes': []})
    for f, v in rec.items():
        if v in (None, ''):
            continue
        if f == 'Dorsais':
            a.setdefault(f, set()).add(v)
        elif a.get(f) in (None, ''):
            a[f] = v  # prevalece o 1.º valor: as fontes são lidas da mais recente para a mais antiga
    if src not in a['Fontes']:
        a['Fontes'].append(src)
    return a

BD_SHEETS = ['SFC | OTC 15Mar', 'SFC | OTC 17Jan', 'SFC | OTC 16Nov', 'SFC | OTC 09Nov']  # mais recente primeiro
for sh in BD_SHEETS:
    for r in rows_of(BD, sh):
        if not r.get('Nome'):
            continue
        upsert(f'BD {sh[-5:]}', {
            'Nome': str(r['Nome']).strip(), 'Género': r.get('Género'), 'Data nascimento': as_date(r.get('Data Nasc')),
            'CC': cc_norm(r.get('CC')), 'Escalão (BD)': r.get('Escalão'), 'Tamanho T-shirt': r.get('Tamanhos'),
            'Tipo camisola': r.get('Tipo'), 'Dorsais': digits(r.get('Dorsal')) or None, 'Email': r.get('Email'),
            'Telemóvel': digits(r.get('Tlm')) or None, 'Observações (BD)': r.get('Observações') if sh.endswith('15Mar') else None,
            'Equipa (BD)': r.get('Equipa'),
        })
for r in rows_of(BD, idx=0):  # lista mestra (Column1..5)
    upsert('BD lista mestra', {'Nome': str(r['Column1']).strip(), 'Género': r['Column2'], 'Data nascimento': as_date(r['Column3']),
                               'Equipa (BD)': r['Column4'], 'CC': cc_norm(r['Column5'])})

def sheet_names(path):
    return openpyxl.load_workbook(path, read_only=True).sheetnames

forms = [r for f in args.form if 'Respostas do Formulário 1' in sheet_names(f) for r in rows_of(f, 'Respostas do Formulário 1')]
seen = set()
for r in sorted(forms, key=lambda x: x['Carimbo de data/hora'], reverse=True):  # resposta mais recente primeiro
    cc = cc_norm(r.get('Cartão  de Cidadão'))
    if (cc, r['Carimbo de data/hora']) in seen:
        continue
    seen.add((cc, r['Carimbo de data/hora']))
    tam = next((v for k, v in r.items() if k.startswith('Tamanho T-shirt')), None)
    a = upsert('Formulário inscrição', {
        'Nome': str(r['Nome completo']).strip(), 'Data nascimento': as_date(r.get('Data de Nascimento')), 'CC': cc,
        'NIF': digits(r.get('NIF')), 'Email': r.get('Email'), 'Morada': r.get('Morada'), 'Tamanho T-shirt': tam,
        'Tipo camisola': r.get('Tipo camisola'), 'Telemóvel': digits(r.get('Telemóvel')) or None,
    })
    a['Nome (formulário)'] = a.get('Nome (formulário)') or str(r['Nome completo']).strip()
    ts = r['Carimbo de data/hora']
    a['Data inscrição (formulário)'] = min(filter(None, [a.get('Data inscrição (formulário)'), ts]))

# Junta registos da mesma pessoa com CC diferente/ausente (gralhas): mesma data de
# nascimento e nome compatível. Prevalece o registo com mais fontes da BD.
recs = sorted(athletes.values(), key=lambda a: -sum(f.startswith('BD') for f in a['Fontes']))
merged = []
for a in recs:
    tgt = next((m for m in merged if m.get('Data nascimento') and m['Data nascimento'] == a.get('Data nascimento')
                and (name_match(m['Nome'], a['Nome']) or name_match(m.get('Nome (formulário)', ''), a['Nome']))), None)
    if not tgt:
        merged.append(a); continue
    for f, v in a.items():
        if f == 'Fontes':
            tgt['Fontes'] += [x for x in v if x not in tgt['Fontes']]
        elif f == 'Dorsais':
            tgt.setdefault('Dorsais', set()).update(v)
        elif f == 'CC' and tgt.get('CC') and v and v != tgt['CC']:
            tgt['CC (alternativo)'] = v
        elif tgt.get(f) in (None, ''):
            tgt[f] = v
    if a.get('Nome') != tgt.get('Nome') and not tgt.get('Nome (formulário)'):
        tgt['Nome (formulário)'] = a['Nome']

# IDs estáveis por ordem alfabética
ordered = sorted(merged, key=lambda a: ' '.join(tokens(a.get('Nome'))))
for i, a in enumerate(ordered, 1):
    a['ID'] = f'SFC-{i:04d}'
by_dorsal = defaultdict(list)
for a in ordered:
    for d in a.get('Dorsais', ()):
        by_dorsal[d].append(a)

# ---------------------------------------------- 2. distâncias por regulamento
G = ['BA', 'BB', 'INF', 'INI', 'JUV', 'JUNF', 'JUNM', 'ADULT']
def D(*v):
    return dict(zip(G, v))
DIST = {  # (época, n.º prova) -> distâncias em metros por grupo (fonte: regulamento de cada prova)
    ('2023/2024', 1): D(200, 600, 1000, 1700, 3000, 3000, 3000, 9000),
    ('2023/2024', 2): D(250, 600, 1000, 1800, 2800, 2800, 7800, 7800),
    ('2023/2024', 4): D(250, 500, 1000, 2000, 3000, 3000, 3000, 4000),
    ('2023/2024', 5): D(250, 400, 800, 1609, 1609, 1609, 1609, 1609),
    ('2023/2024', 6): D(200, 600, 1000, 1700, 3000, 3000, 3000, 9000),
    ('2023/2024', 7): D(250, 600, 1000, 2000, 3950, 3950, 7100, 7100),
    ('2024/2025', 1): D(200, 600, 1000, 1700, 3000, 3000, 3000, 9000),
    ('2024/2025', 2): D(250, 590, 1000, 1700, 3700, 3700, 5850, 5850),
    ('2024/2025', 3): D(250, 500, 1000, 2000, 3000, 3000, 3000, None),  # adultos: 3 ou 4 km (ver abaixo)
    ('2024/2025', 4): D(250, 600, 1000, 1800, 2800, 2800, 7800, 7800),
    ('2024/2025', 5): D(250, 600, 1000, 2000, None, None, 7100, 7100),  # juvenis/juniores F: 3950 ou 4000
    ('2024/2025', 7): D(250, 400, 800, 1609, 1609, 1609, 1609, 1609),
    ('2024/2025', 8): D(250, 500, 1000, 2000, 3000, 3000, 3000, 3500),
    ('2024/2025', 9): D(200, 600, 1000, 1700, 3000, 3000, 3000, 9000),
    ('2025/2026', 1): D(250, 590, 1000, 1700, 3700, 3700, 5850, 5850),
    ('2025/2026', 2): D(200, 600, 1000, 1700, 3000, 3000, 3000, 8000),
    ('2025/2026', 3): D(200, 400, 600, 1000, 1600, 1600, 1600, 1600),
    ('2025/2026', 4): D(250, 600, 1000, 1800, 2800, 2800, 7800, 7800),
    ('2025/2026', 6): D(250, 400, 800, 1609, 1609, 1609, 1609, 1609),
    ('2025/2026', 7): D(250, 500, 1000, 2000, 3000, 3000, 3000, 3500),
    ('2025/2026', 8): D(200, 600, 1000, 1700, 3000, 3000, 7000, 7000),
}
REG = {}
for line in open(args.regs, encoding='utf-8'):
    if line.strip():
        k, u = line.strip().split('\t')
        REG[k] = u
def reg_url(season, n):
    y = season[2:4] + '-' + season[7:9]
    return REG.get(f'{n}a-prova-{y}') or REG.get(f'{n}-{y}', '')

VET_ROMAN = {'I': 1, 'II': 2, 'III': 3, 'IV': 4, 'V': 5, 'VI': 6, 'VII': 7, 'VIII': 8}
def group(esc, sexo):
    return {'Benjamins A': 'BA', 'Benjamins B': 'BB', 'Infantis': 'INF', 'Iniciados': 'INI', 'Juvenis': 'JUV'}.get(esc) or \
        ('JUNF' if esc == 'Juniores' and sexo == 'F' else 'JUNM' if esc == 'Juniores' else 'ADULT')

def distance(season, n, esc, sexo):
    d = DIST.get((season, n))
    if not d:
        return None, 'sem regulamento'
    g = group(esc, sexo)
    if (season, n) == ('2024/2025', 3) and g == 'ADULT':
        vet = re.search(r'(I|II|III|IV|V|VI|VII|VIII)$', esc)
        lvl = VET_ROMAN[vet.group(1)] if vet else 0
        return (4000 if sexo == 'M' and lvl <= 3 else 3000), ''
    if (season, n) == ('2024/2025', 5) and g in ('JUV', 'JUNF'):
        return 3950, 'regulamento indica 3950 m e 4000 m'
    return d[g], ''

# ------------------------------------------------------- 3. resultados do troféu
with open(RESULTS, encoding='utf-8-sig') as f:
    results = list(csv.DictReader(f, delimiter=';'))

def find_athlete(r):
    cands = []
    if r['Época'] == '2025/2026' and r['Dorsal'] in by_dorsal:
        cands = [a for a in by_dorsal[r['Dorsal']] if name_match(a['Nome'], r['Nome']) or name_match(a.get('Nome (formulário)', ''), r['Nome'])]
        if len(cands) == 1:
            return cands[0], 'dorsal + nome'
    yr = r['Ano nascimento']
    cands = [a for a in ordered if a.get('Data nascimento') and str(a['Data nascimento'].year) == yr
             and (name_match(a['Nome'], r['Nome']) or name_match(a.get('Nome (formulário)', ''), r['Nome']))]
    if len(cands) == 1:
        return cands[0], 'nome + ano nasc.'
    return None, ('ambíguo' if cands else 'não encontrado na BD')

ext_ids, unmatched = {}, []
for r in results:
    a, how = find_athlete(r)
    if not a:  # atleta sem registo na BD: ID externo estável por nome + ano
        k = (' '.join(tokens(r['Nome'])), r['Ano nascimento'])
        ext_ids.setdefault(k, f'EXT-{len(ext_ids) + 1:04d}')
        r['ID atleta'] = ext_ids[k]
        unmatched.append((r, how))
    else:
        r['ID atleta'] = a['ID']
    r['Associação'] = how
    dist, note = distance(r['Época'], int(r['N.º prova']), r['Escalão (base)'], r['Sexo'])
    r['Distância (m)'] = dist
    r['Nota distância'] = note
    sec = float(r['Tempo (s)']) if r['Tempo (s)'] else None
    r['Ritmo (min/km)'] = ''
    if sec and dist:
        pace = sec / (dist / 1000)
        r['Ritmo (min/km)'] = f'{int(pace // 60)}:{int(pace % 60):02d}'
        if pace < 150 or pace > 900:  # < 2:30/km ou > 15:00/km: distância real provavelmente diferente
            r['Nota distância'] = (note + '; ' if note else '') + 'ritmo implausível — distância real pode diferir do regulamento'

# --------------------------------------------- 4. resumo por atleta (BD)
res_by = defaultdict(list)
for r in results:
    res_by[r['ID atleta']].append(r)
for a in ordered:
    rs = res_by.get(a['ID'], [])
    a['N.º provas troféu (total)'] = len(rs)
    for s in ('2023/2024', '2024/2025', '2025/2026'):
        a[f'Provas {s[2:4]}/{s[7:9]}'] = sum(1 for r in rs if r['Época'] == s)
    a['Pontos troféu 25/26'] = sum(int(r['Pontos troféu']) for r in rs if r['Época'] == '2025/2026' and r['Pontos troféu'].isdigit())
    best = [int(r['Classificação no escalão']) for r in rs if r['Classificação no escalão'].isdigit()]
    a['Melhor classificação no escalão'] = min(best) if best else None
    b = a.get('Data nascimento')
    a['Idade'] = (REF_DATE.year - b.year - ((REF_DATE.month, REF_DATE.day) < (b.month, b.day))) if b else None
    a['Dorsal 25/26'] = ', '.join(sorted(a.get('Dorsais', ())))
    a['Fontes'] = ', '.join(a['Fontes'])

# --------------------------------------------- 5. comparação entre épocas
SEAS = ['2023/2024', '2024/2025', '2025/2026']
groups = defaultdict(dict)
for r in results:
    groups[(r['ID atleta'], r['Prova (nome base)'])][r['Época']] = r
comparison = []
for (aid, prova), by in groups.items():
    any_r = next(iter(by.values()))
    line = {'ID atleta': aid, 'Nome': any_r['Nome'], 'Ano nascimento': any_r['Ano nascimento'], 'Sexo': any_r['Sexo'],
            'Prova (nome base)': prova, 'N.º épocas': len(by)}
    for s in SEAS:
        y = s[2:4] + '/' + s[7:9]
        r = by.get(s)
        line[f'Escalão {y}'] = r['Escalão (base)'] if r else None
        line[f'Distância {y} (m)'] = r['Distância (m)'] if r else None
        line[f'Class. {y}'] = int(r['Classificação no escalão']) if r and r['Classificação no escalão'].isdigit() else None
        line[f'Tempo {y}'] = r['Tempo / Marca'] if r else None
        line[f'Ritmo {y}'] = r['Ritmo (min/km)'] if r else None
    timed = [(s, by[s]) for s in SEAS if s in by and by[s]['Tempo (s)']]
    if len(timed) >= 2:
        (_, p), (_, l) = timed[-2], timed[-1]
        same = p['Distância (m)'] and p['Distância (m)'] == l['Distância (m)']
        line['Mesma distância'] = 'Sim' if same else 'Não'
        line['Evolução tempo (s)'] = round(float(l['Tempo (s)']) - float(p['Tempo (s)']), 2) if same else None
        if p['Distância (m)'] and l['Distância (m)']:
            pp = float(p['Tempo (s)']) / (p['Distância (m)'] / 1000)
            lp = float(l['Tempo (s)']) / (l['Distância (m)'] / 1000)
            line['Evolução ritmo (s/km)'] = round(lp - pp, 1)
    comparison.append(line)
comparison.sort(key=lambda l: (-l['N.º épocas'], l['Prova (nome base)'], l['Nome']))

# --------------------------------------------- 6. inscrições Corrida de Natal (Formulário 2)
natal = []
for r in (r for f in args.form if 'Respostas do Formulário 2' in sheet_names(f) for r in rows_of(f, 'Respostas do Formulário 2')):
    full = f"{r['Nome']} {r['Apelido']}".strip()
    tel = digits(r.get('Telemóvel'))
    c = [a for a in ordered if (tel and a.get('Telemóvel') == tel) or name_match(a['Nome'], full) or name_match(a.get('Nome (formulário)', ''), full)]
    tam = next((v for k, v in r.items() if k.startswith('Tamanho')), None)
    dist_ = next((v for k, v in r.items() if k.startswith('Distância')), None)
    natal.append({'Data inscrição': r['Carimbo de data/hora'], 'Nome': full, 'Telemóvel': tel, 'Tamanho T-shirt': tam,
                  'Distância': dist_, 'ID atleta': c[0]['ID'] if len(c) == 1 else None,
                  'Associação': 'telemóvel/nome' if len(c) == 1 else ('ambíguo' if c else 'não encontrado na BD')})

# --------------------------------------------- 7. escrever o Excel
wb = openpyxl.Workbook()
HEAD = PatternFill('solid', fgColor='004A8E')
def sheet(title, fields, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(fields)
    for r in rows:
        ws.append([(', '.join(sorted(v)) if isinstance(v, set) else v) for v in (r.get(f) for f in fields)])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    for i, f in enumerate(fields, 1):
        vals = [len(str(f))] + [len(str(x.get(f) or '')) for x in rows[:300]]
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(8, max(vals) + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (dt.date, dt.datetime)):
                c.number_format = 'dd/mm/yyyy' if not isinstance(c.value, dt.datetime) else 'dd/mm/yyyy hh:mm'
    return ws

A_FIELDS = ['ID', 'Nome', 'Nome (formulário)', 'Género', 'Data nascimento', 'Idade', 'Escalão (BD)', 'CC', 'CC (alternativo)', 'NIF', 'Email', 'Telemóvel',
            'Morada', 'Tamanho T-shirt', 'Tipo camisola', 'Dorsal 25/26', 'Equipa (BD)', 'Observações (BD)', 'Data inscrição (formulário)',
            'N.º provas troféu (total)', 'Provas 23/24', 'Provas 24/25', 'Provas 25/26', 'Pontos troféu 25/26',
            'Melhor classificação no escalão', 'Fontes']
sheet('Atletas', A_FIELDS, ordered, first=True)

R_FIELDS = ['ID atleta', 'Nome', 'Época', 'N.º prova', 'Prova', 'Prova (nome base)', 'Edição', 'Data', 'Escalão', 'Escalão (base)',
            'Sexo', 'Classificação no escalão', 'Dorsal', 'Ano nascimento', 'Tempo / Marca', 'Tempo (s)', 'Distância (m)',
            'Ritmo (min/km)', 'Pontos equipas', 'Pontos troféu', 'Nota distância', 'Associação', 'Fonte']
for r in results:
    for f in ('N.º prova', 'Classificação no escalão', 'Pontos equipas', 'Pontos troféu', 'Ano nascimento'):
        r[f] = int(r[f]) if str(r[f]).isdigit() else (r[f] or None)
    r['Tempo (s)'] = float(r['Tempo (s)']) if r['Tempo (s)'] else None
    r['Data'] = dt.date.fromisoformat(r['Data'])
sheet('Resultados', R_FIELDS, sorted(results, key=lambda r: (r['Data'], r['Escalão (base)'], r['Classificação no escalão'] or 999)))

C_FIELDS = ['ID atleta', 'Nome', 'Ano nascimento', 'Sexo', 'Prova (nome base)', 'N.º épocas']
for s in SEAS:
    y = s[2:4] + '/' + s[7:9]
    C_FIELDS += [f'Escalão {y}', f'Distância {y} (m)', f'Class. {y}', f'Tempo {y}', f'Ritmo {y}']
C_FIELDS += ['Mesma distância', 'Evolução tempo (s)', 'Evolução ritmo (s/km)']
sheet('Comparação épocas', C_FIELDS, comparison)

dist_rows = []
labels = {'BA': 'Benjamins A', 'BB': 'Benjamins B', 'INF': 'Infantis', 'INI': 'Iniciados', 'JUV': 'Juvenis',
          'JUNF': 'Juniores F', 'JUNM': 'Juniores M', 'ADULT': 'Seniores e Veteranos'}
prova_name = {(r['Época'], r['N.º prova']): (r['Prova'], r['Prova (nome base)']) for r in results}
for (s, n), d in sorted(DIST.items()):
    pn = prova_name.get((s, n), ('', ''))
    row = {'Época': s, 'N.º prova': n, 'Prova': pn[0], 'Prova (nome base)': pn[1], 'Regulamento': reg_url(s, n)}
    row.update({labels[g]: d[g] for g in G})
    if (s, n) == ('2024/2025', 3):
        row['Seniores e Veteranos'] = 'Sen./Vet. I-III M: 4000; restantes: 3000'
    if (s, n) == ('2024/2025', 5):
        row['Juvenis'] = row['Juniores F'] = '3950 / 4000'
    dist_rows.append(row)
sheet('Distâncias', ['Época', 'N.º prova', 'Prova', 'Prova (nome base)'] + list(labels.values()) + ['Regulamento'], dist_rows)

sheet('Corrida Natal (Form. 2)', ['Data inscrição', 'Nome', 'Telemóvel', 'Tamanho T-shirt', 'Distância', 'ID atleta', 'Associação'], natal)

un_rows = [{'ID atleta': r['ID atleta'], 'Nome': r['Nome'], 'Ano nascimento': r['Ano nascimento'], 'Época': r['Época'],
            'Prova': r['Prova'], 'Dorsal': r['Dorsal'], 'Motivo': how} for r, how in unmatched]
sheet('Por associar', ['ID atleta', 'Nome', 'Ano nascimento', 'Época', 'Prova', 'Dorsal', 'Motivo'], un_rows)

notes = wb.create_sheet('Notas')
for line in [
    'Ficheiro consolidado Serrado FC — atletas e resultados do Troféu Almada em Atletismo "Mário Pinto Claro"',
    f'Gerado em {dt.date.today():%d/%m/%Y}. CONTÉM DADOS PESSOAIS (CC, NIF, morada, contactos, dados de menores): uso interno, não partilhar.',
    '',
    'Fontes:',
    '• BD_20260315 - Charneca.xlsx: lista mestra + inscrições por prova (09Nov, 16Nov, 17Jan, 15Mar). A folha mais recente prevalece.',
    '• Formulários de inscrição Troféu de Almada 2025 (Respostas 1 e 2): NIF, morada, t-shirt; Formulário 2 = Corrida de Natal (Leblon Run Club).',
    '• Resultados oficiais: tatletismo-almada.pt (tabelas de cada prova 23/24 e 24/25; CSV publicados 25/26).',
    '• Distâncias: regulamento de cada prova (link na folha Distâncias).',
    '',
    'Regras de associação:',
    '• Atletas da BD agrupados pelo n.º de CC (ou nome + data de nascimento quando não há CC).',
    '• Resultados 25/26 associados pelo dorsal do troféu + nome; épocas anteriores por nome + ano de nascimento.',
    '• Resultados sem correspondência na BD recebem um ID "EXT-" (ver folha Por associar).',
    '',
    'Atenção na comparação:',
    '• Só compare tempos com "Mesma distância = Sim". Ex.: Corrida Egas Moniz teve 9 km (23/24, 24/25) e 7 km (25/26).',
    '• "Evolução ritmo (s/km)" permite comparar mesmo com distâncias diferentes (negativo = mais rápido).',
    '• Nos escalões jovens a distância muda quando o atleta sobe de escalão.',
    '• "Ritmo implausível" assinala provas em que a distância real parece diferente da do regulamento (frequente nos Benjamins).',
]:
    notes.append([line])
notes.column_dimensions['A'].width = 140
wb.move_sheet('Notas', offset=-(len(wb.sheetnames) - 1))
wb.save(OUT)

# --------------------------------------------- 8. CSV de importação (results.csv → backoffice; athletes.csv → tabela athletes, ver services/README.md)
if args.import_dir:
    os.makedirs(args.import_dir, exist_ok=True)
    with open(os.path.join(args.import_dir, 'athletes.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['code', 'name', 'gender', 'birth_date', 'id_number', 'tax_number', 'email', 'phone', 'address', 'sport_slug', 'category', 'shirt_size'])
        for a in ordered:
            w.writerow([a['ID'], a.get('Nome'), a.get('Género'), a.get('Data nascimento') or '', a.get('CC'), a.get('NIF'), a.get('Email'),
                        a.get('Telemóvel'), a.get('Morada'), 'atletismo', a.get('Escalão (BD)'), a.get('Tamanho T-shirt')])
    with open(os.path.join(args.import_dir, 'results.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['athlete_code', 'athlete_name', 'birth_year', 'season', 'round', 'race', 'race_base', 'race_date', 'category',
                    'place', 'bib', 'time', 'time_s', 'distance_m', 'trophy_points', 'team_points', 'source_url'])
        for r in results:
            w.writerow([r['ID atleta'] if r['ID atleta'].startswith('SFC-') else '', r['Nome'], r['Ano nascimento'], r['Época'], r['N.º prova'],
                        r['Prova'], r['Prova (nome base)'], r['Data'], r['Escalão (base)'], r['Classificação no escalão'] or '', r['Dorsal'],
                        r['Tempo / Marca'], r['Tempo (s)'] or '', r['Distância (m)'] or '', r['Pontos troféu'] or '', r['Pontos equipas'] or '', r['Fonte']])
    print(f'CSV de importação → {args.import_dir}/athletes.csv, results.csv')

# --------------------------------------------- resumo (sem dados pessoais)
m = Counter(r['Associação'] for r in results)
print(f'Atletas na BD consolidada: {len(ordered)}  (com resultados no troféu: {sum(1 for a in ordered if a["N.º provas troféu (total)"])})')
print('Resultados:', len(results), dict(m))
print('Atletas externos (EXT):', len(ext_ids))
print('Comparações:', len(comparison), '| com 2+ épocas:', sum(1 for c in comparison if c['N.º épocas'] > 1),
      '| mesma distância:', sum(1 for c in comparison if c.get('Mesma distância') == 'Sim'))
print('Corrida Natal:', len(natal), dict(Counter(n['Associação'] for n in natal)))
print('Avisos de ritmo:', sum(1 for r in results if 'implausível' in (r['Nota distância'] or '')),
      Counter(r['Escalão (base)'] for r in results if 'implausível' in (r['Nota distância'] or '')).most_common(6))
print('Fontes BD:', Counter(a['Fontes'] for a in ordered).most_common(8))
